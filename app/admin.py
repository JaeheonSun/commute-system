from flask import Blueprint, render_template, redirect, url_for, request, flash
from flask_login import login_required
from werkzeug.security import generate_password_hash
from .models import User, WorkLog
from . import db
from .utils import admin_required
import json
from datetime import datetime, timedelta, time
from sqlalchemy import extract, func, text
from .services import calculate_work_time, now_kst, is_weekend, is_holiday
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from flask import send_file

admin_bp = Blueprint("admin", __name__, url_prefix="/admin")

WORK_TYPE_LABELS = {
    "office": "재단근무",
    "remote": "재택근무",
    "flexible": "시차(유연근무)",
    "half_day": "반차",
    "annual_leave": "연차",
    "special_leave": "특별휴가",
    "substitute_leave": "대체휴무",
}

@admin_bp.route("/")
@login_required
@admin_required
def dashboard():
    users = User.query.all()
    return render_template("admin_dashboard.html", users=users)


@admin_bp.route("/create-user", methods=["POST"])
@login_required
@admin_required
def create_user():
    username = request.form["username"]
    password = request.form["password"]
    role = request.form["role"]

    if User.query.filter_by(username=username).first():
        flash("이미 존재하는 username입니다.", "warning")
        return redirect(url_for("admin.dashboard"))

    user = User(
        username=username,
        password_hash=generate_password_hash(password),
        role=role
    )
    db.session.add(user)
    db.session.commit()

    return redirect(url_for("admin.dashboard"))


@admin_bp.route("/work-summary")
@login_required
@admin_required
def work_summary():
    year = request.args.get("year", type=int)
    month = request.args.get("month", type=int)
    selected_user_id = request.args.get("user_id", type=int)
    day = request.args.get("day", type=int)

    if not year:
        year = now_kst().year

    years = (
        db.session.query(extract("year", WorkLog.work_date))
        .distinct()
        .order_by(extract("year", WorkLog.work_date).desc())
        .all()
    )
    years = [int(y[0]) for y in years]
    if year not in years:
        years.insert(0, year)

    all_users = User.query.order_by(User.username).all()

    use_detail = bool(selected_user_id or day)

    if use_detail:
        q = db.session.query(WorkLog).filter(
            extract("year", WorkLog.work_date) == year
        )
        if month:
            q = q.filter(extract("month", WorkLog.work_date) == month)
        if day:
            q = q.filter(extract("day", WorkLog.work_date) == day)
        if selected_user_id:
            q = q.filter(WorkLog.user_id == selected_user_id)
        records = q.order_by(WorkLog.work_date, WorkLog.username).all()
        results = None
    else:
        q = (
            db.session.query(
                User.username,
                func.sum(WorkLog.total_minutes).label("total_minutes"),
                func.sum(WorkLog.overtime_minutes).label("overtime_minutes"),
            )
            .join(WorkLog, WorkLog.user_id == User.id)
            .filter(extract("year", WorkLog.work_date) == year)
        )
        if month:
            q = q.filter(extract("month", WorkLog.work_date) == month)
        results = q.group_by(User.id).all()
        records = None

    return render_template(
        "admin_work_summary.html",
        years=years,
        all_users=all_users,
        selected_year=year,
        selected_month=month,
        selected_user_id=selected_user_id,
        selected_day=day,
        use_detail=use_detail,
        results=results,
        records=records,
        work_type_labels=WORK_TYPE_LABELS,
    )


@admin_bp.route("/export")
@login_required
@admin_required
def export_worklog():
    year = request.args.get("year", type=int)
    month = request.args.get("month", type=int)
    if not year:
        year = now_kst().year

    thin = Side(style="thin")
    full_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_middle = Alignment(horizontal="left", vertical="center")
    header_font = Font(bold=True, size=10)
    header_fill = PatternFill(fill_type="solid", fgColor="BDD7EE")

    # Column widths matching example_format.xlsx
    COL_WIDTHS = [3.22, 4.44, 12.55, 12.55, 13.66, 10.89, 10.89, 10.89, 10.89, 9.22, 9.22, 9.22]

    def fmt_min(minutes):
        m = minutes or 0
        return f"{m // 60:02d}:{m % 60:02d}"

    def apply_border_row(ws, row, col_start, col_end):
        for col in range(col_start, col_end + 1):
            ws.cell(row=row, column=col).border = full_border

    wb = Workbook()
    default_ws = wb.active

    users = User.query.all()
    any_created = False

    for user in users:
        q = WorkLog.query.filter(
            WorkLog.user_id == user.id,
            extract("year", WorkLog.work_date) == year,
        )
        if month:
            q = q.filter(extract("month", WorkLog.work_date) == month)
        records = q.order_by(WorkLog.work_date).all()
        if not records:
            continue

        ws = wb.create_sheet(title=user.username[:31])
        any_created = True

        # Column widths
        for i, w in enumerate(COL_WIDTHS, start=1):
            ws.column_dimensions[chr(64 + i)].width = w

        # Row 1: 출퇴근현황표 (A1:L1 merged)
        ws.merge_cells("A1:L1")
        c = ws["A1"]
        c.value = f"{month}월 출퇴근현황표" if month else f"{year}년 출퇴근현황표"
        c.font = Font(bold=True, size=14)
        c.alignment = center
        ws.row_dimensions[1].height = 30

        # Row 2: headers (A2:B2 merged for date)
        ws.merge_cells("A2:B2")
        header_map = [
            (1, "월   일"),
            (3, "근무형태"),
            (4, "내용"),
            (5, "성        명"),
            (6, "퇴근시간"),
            (7, "시간외근무"),
            (8, "시간외근무\n총시간"),
            (9, "확인"),
            (10, "결재1"),
            (11, "결재2"),
            (12, "비      고"),
        ]
        for col, val in header_map:
            c = ws.cell(row=2, column=col, value=val)
            c.font = header_font
            c.alignment = center
            c.fill = header_fill
        apply_border_row(ws, 2, 1, 12)
        ws.row_dimensions[2].height = 32.25

        # Data rows starting at row 3
        cumulative_ot = 0
        for ridx, r in enumerate(records, start=3):
            cumulative_ot += r.overtime_minutes or 0
            row_vals = [
                r.work_date.month,
                r.work_date.day,
                WORK_TYPE_LABELS.get(r.work_type, r.work_type or ""),
                r.start_time.strftime("%H:%M") if r.start_time else "",
                user.username,
                r.end_time.strftime("%H:%M") if r.end_time else "",
                fmt_min(r.overtime_minutes or 0),
                fmt_min(cumulative_ot),
                "", "", "",
                r.remarks or "",
            ]
            for cidx, val in enumerate(row_vals, start=1):
                c = ws.cell(row=ridx, column=cidx, value=val)
                c.border = full_border
                c.alignment = left_middle if cidx == 12 else center
            ws.row_dimensions[ridx].height = 19.5

        # Footer
        last_data_row = len(records) + 2
        ws.cell(row=last_data_row + 2, column=9, value="출력일")
        ws.cell(row=last_data_row + 2, column=10, value=now_kst().strftime("%Y-%m-%d"))
        ws.cell(row=last_data_row + 3, column=7, value="재단법인")
        ws.cell(row=last_data_row + 3, column=8, value="한사람")

    if any_created:
        wb.remove(default_ws)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    filename = f"worklog_{year}{('_%02d' % month) if month else ''}.xlsx"
    return send_file(bio, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@admin_bp.route("/worklog/add", methods=["GET", "POST"])
@login_required
@admin_required
def add_worklog():
    users = User.query.all()

    if request.method == "POST":
        user_id = request.form["user_id"]
        user = User.query.get(int(user_id))
        work_date = datetime.strptime(request.form["work_date"], "%Y-%m-%d").date()

        work_type = request.form.get("work_type", "office")
        remarks = request.form.get("remarks", "")

        existing = WorkLog.query.filter_by(
            user_id=user_id,
            work_date=work_date
        ).first()

        if existing:
            flash("이미 해당 날짜의 근무 기록이 존재합니다.", "error")
            return redirect(request.url)

        # Leaves and substitute: no times required
        if work_type in ("annual_leave", "special_leave", "substitute_leave"):
            if work_type == "substitute_leave" and not remarks.strip():
                flash("대체휴무인 경우 비고(사유)를 입력해야 합니다.", "warning")
                return redirect(request.url)

            worklog = WorkLog(
                user_id=user_id,
                username=user.username,
                work_date=work_date,
                start_time=None,
                end_time=None,
                total_minutes=0,
                overtime_minutes=0,
                work_type=work_type,
                remarks=remarks
            )

            db.session.add(worklog)
            db.session.commit()

            flash("근무 기록이 추가되었습니다.", "success")
            return redirect(url_for("admin.dashboard"))

        # Expect start/end for other types
        try:
            start_t = datetime.strptime(request.form["start_time"], "%H:%M").time()
            end_t = datetime.strptime(request.form["end_time"], "%H:%M").time()
        except Exception:
            flash("출근/퇴근 시간을 정확히 입력해주세요.", "warning")
            return redirect(request.url)

        start_dt = datetime.combine(work_date, start_t)
        end_dt = datetime.combine(work_date, end_t)

        if end_dt <= start_dt:
            flash("퇴근 시간은 출근 시간보다 이후여야 합니다.", "warning")
            return redirect(request.url)

        worklog = WorkLog(
            user_id=user_id,
            username = user.username,
            work_date=work_date,
            start_time=start_t,
            end_time=end_t,
            total_minutes=0,
            overtime_minutes=0,
            work_type=work_type,
            remarks=remarks
        )

        calculate_work_time(worklog)

        db.session.add(worklog)
        db.session.commit()

        flash("근무 기록이 추가되었습니다.", "success")
        return redirect(url_for("admin.dashboard"))

    return render_template(
        "admin_worklog_form.html",
        users=users
    )



@admin_bp.route("/worklog/edit", methods=["GET", "POST"])
@login_required
@admin_required
def edit_worklog_by_query():
    users = User.query.all()
    record = None
    selected_user = None
    selected_date = None

    if request.method == "POST":
        selected_user = int(request.form["user_id"])
        selected_date = datetime.strptime(
            request.form["work_date"], "%Y-%m-%d"
        ).date()

        # 🔍 STEP 1: 근무 기록 조회
        record = WorkLog.query.filter_by(
            user_id=selected_user,
            work_date=selected_date
        ).first()

        if not record:
            flash("해당 날짜의 근무 기록이 없습니다.", "error")
            return render_template(
                "admin_worklog_edit.html",
                users=users,
                record=None,
                selected_user=selected_user,
                selected_date=selected_date
            )

        # ✏️ STEP 2: 수정 요청인지 확인
        if "action" in request.form or "work_type" in request.form:
            # Update work_type and remarks first
            new_work_type = request.form.get("work_type", record.work_type)
            new_remarks = request.form.get("remarks", record.remarks or "")

            record.work_type = new_work_type
            record.remarks = new_remarks

            if new_work_type in ("annual_leave", "special_leave", "substitute_leave"):
                if new_work_type == "substitute_leave" and not new_remarks.strip():
                    flash("대체휴무인 경우 비고(사유)를 입력해야 합니다.", "warning")
                    return redirect(request.url)

                record.start_time = None
                record.end_time = None
                record.total_minutes = 0
                record.overtime_minutes = 0
                db.session.commit()
                flash("근무 기록이 수정되었습니다.", "success")
                return redirect(url_for("admin.dashboard"))

            # For other types, expect start/end times
            try:
                start_time = datetime.strptime(request.form["start_time"], "%H:%M").time()
                end_time = datetime.strptime(request.form["end_time"], "%H:%M").time()
            except Exception:
                flash("출근/퇴근 시간을 정확히 입력해주세요.", "warning")
                return redirect(request.url)

            start_dt = datetime.combine(selected_date, start_time)
            end_dt = datetime.combine(selected_date, end_time)

            if end_dt <= start_dt:
                flash("퇴근 시간은 출근 시간보다 이후여야 합니다.", "warning")
                return redirect(request.url)

            record.start_time = start_time
            record.end_time = end_time

            calculate_work_time(record)
            db.session.commit()

            flash("근무 기록이 수정되었습니다.", "success")
            return redirect(url_for("admin.dashboard"))

    return render_template(
        "admin_worklog_edit.html",
        users=users,
        record=record,
        selected_user=selected_user,
        selected_date=selected_date
    )


@login_required
@admin_required
@admin_bp.route("/worklog/delete", methods=["POST"])
def delete_worklog():
    user_id = request.form.get("user_id")
    work_date = datetime.strptime(
        request.form.get("work_date"),
        "%Y-%m-%d"
    ).date()

    record = WorkLog.query.filter_by(
        user_id=user_id,
        work_date=work_date
    ).first()

    if not record:
        flash("삭제할 근무 기록이 존재하지 않습니다.", "danger")
        return redirect(url_for("admin.edit_worklog_by_query"))

    db.session.delete(record)
    db.session.commit()

    flash("근무 기록이 삭제되었습니다.", "success")
    return redirect(url_for("admin.edit_worklog_by_query"))



@admin_bp.route("/user/update", methods=["POST"])
@login_required
@admin_required
def update_user():
    user_id = request.form["user_id"]
    new_username = request.form.get("username")
    new_password = request.form.get("password")

    user = User.query.get_or_404(user_id)

    # 관리자 username 변경 방지하고 싶으면 여기서 체크 가능
    if new_username:
        existing = User.query.filter(
            User.username == new_username,
            User.id != user.id
        ).first()
        if existing:
            flash("이미 존재하는 username입니다.", "warning")
            return redirect(url_for("admin.dashboard"))

        user.username = new_username

    if new_password:
        user.password_hash = generate_password_hash(new_password)

    db.session.commit()
    flash("유저 정보가 수정되었습니다.", "success")
    return redirect(url_for("admin.dashboard"))


@admin_bp.route("/user/delete/<int:user_id>", methods=["POST"])
@login_required
@admin_required
def delete_user(user_id):
    user = User.query.get_or_404(user_id)

    if user.role == "admin":
        flash("관리자 계정은 삭제할 수 없습니다.", "error")
        return redirect(url_for("admin.dashboard"))

    db.session.delete(user)
    db.session.commit()

    flash("유저가 삭제되었습니다.", "success")
    return redirect(url_for("admin.dashboard"))


# ─────────────────────────────────────────────
# 휴가 관리
# ─────────────────────────────────────────────

LEAVE_TYPES = ("annual_leave", "special_leave", "substitute_leave")


def _parse_leave_dates(form):
    """폼에서 날짜 목록 파싱. 실패 시 ValueError."""
    mode = form.get("mode", "single")
    if mode == "range":
        start = datetime.strptime(form.get("start_date", ""), "%Y-%m-%d").date()
        end   = datetime.strptime(form.get("end_date",   ""), "%Y-%m-%d").date()
        if end < start:
            raise ValueError("종료일은 시작일보다 이후여야 합니다.")
        dates, cur = [], start
        while cur <= end:
            dates.append(cur)
            cur += timedelta(days=1)
        return dates
    else:
        return [datetime.strptime(form.get("leave_date", ""), "%Y-%m-%d").date()]


def _resolve_users(form):
    """폼에서 대상 직원 목록 반환. 선택 없으면 None."""
    if form.get("select_all") == "1":
        return User.query.filter(User.username != "peabodu509").all()
    ids = form.getlist("user_ids")
    if not ids:
        return None
    return User.query.filter(User.id.in_([int(i) for i in ids])).all()


@admin_bp.route("/leave-management")
@login_required
@admin_required
def leave_management():
    users = User.query.filter(User.username != "peabodu509").order_by(User.username).all()
    cur_year = now_kst().year

    records = WorkLog.query.filter(
        WorkLog.work_type.in_(LEAVE_TYPES),
        extract("year", WorkLog.work_date).in_([cur_year - 1, cur_year, cur_year + 1]),
    ).all()

    color_map = {
        "annual_leave":    "#f39c12",
        "special_leave":   "#2980b9",
        "substitute_leave":"#8e44ad",
    }
    events = [
        {
            "title": r.username,
            "start": r.work_date.isoformat(),
            "color": color_map.get(r.work_type, "#aaa"),
            "extendedProps": {
                "username":  r.username,
                "work_type": WORK_TYPE_LABELS.get(r.work_type, r.work_type),
                "remarks":   r.remarks or "",
            },
        }
        for r in records
    ]

    return render_template(
        "admin_leave_management.html",
        users=users,
        events_json=json.dumps(events, ensure_ascii=False),
        work_type_labels=WORK_TYPE_LABELS,
    )


@admin_bp.route("/leave-management/add", methods=["POST"])
@login_required
@admin_required
def add_leave():
    users = _resolve_users(request.form)
    if not users:
        flash("직원을 한 명 이상 선택해주세요.", "warning")
        return redirect(url_for("admin.leave_management"))

    try:
        dates = _parse_leave_dates(request.form)
    except ValueError as e:
        flash(str(e) if str(e) else "날짜를 올바르게 입력해주세요.", "warning")
        return redirect(url_for("admin.leave_management"))

    invalid = [d for d in dates if is_weekend(d) or is_holiday(d)]
    if invalid:
        flash(
            "주말 또는 공휴일이 포함되어 있어 등록할 수 없습니다: "
            + ", ".join(d.strftime("%Y-%m-%d") for d in invalid),
            "warning",
        )
        return redirect(url_for("admin.leave_management"))

    work_type = request.form.get("work_type", "annual_leave")
    remarks   = request.form.get("remarks", "")
    added = skipped = 0

    for user in users:
        for d in dates:
            if WorkLog.query.filter_by(user_id=user.id, work_date=d).first():
                skipped += 1
                continue
            db.session.add(WorkLog(
                user_id=user.id,
                username=user.username,
                work_date=d,
                start_time=time(8, 0),
                end_time=time(8, 0),
                total_minutes=0,
                overtime_minutes=0,
                work_type=work_type,
                remarks=remarks,
            ))
            added += 1

    db.session.commit()
    if added:
        flash(f"휴가 기록 {added}건이 등록되었습니다.", "success")
    if skipped:
        flash(f"이미 기록이 있어 건너뛴 건수: {skipped}건", "warning")
    return redirect(url_for("admin.leave_management"))


@admin_bp.route("/leave-management/cancel", methods=["POST"])
@login_required
@admin_required
def cancel_leave():
    users = _resolve_users(request.form)
    if not users:
        flash("직원을 한 명 이상 선택해주세요.", "warning")
        return redirect(url_for("admin.leave_management"))

    try:
        dates = _parse_leave_dates(request.form)
    except ValueError as e:
        flash(str(e) if str(e) else "날짜를 올바르게 입력해주세요.", "warning")
        return redirect(url_for("admin.leave_management"))

    deleted = 0
    for user in users:
        for d in dates:
            record = WorkLog.query.filter(
                WorkLog.user_id == user.id,
                WorkLog.work_date == d,
                WorkLog.work_type.in_(LEAVE_TYPES),
            ).first()
            if record:
                db.session.delete(record)
                deleted += 1

    db.session.commit()
    if deleted:
        flash(f"휴가 기록 {deleted}건이 취소되었습니다.", "success")
    else:
        flash("취소할 휴가 기록이 없습니다.", "warning")
    return redirect(url_for("admin.leave_management"))
